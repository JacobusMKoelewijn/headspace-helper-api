# import warnings
import os
import openpyxl
from openpyxl.chart import ScatterChart, Reference, Series
from openpyxl.chart.trendline import Trendline, TrendlineLabel
from openpyxl.chart.data_source import NumFmt
from openpyxl.chart.text import RichText
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font
from . import root_dir
from . import cells_with_reference
from .data import Solvent, Diluent, Sample
import tempfile

# warnings.simplefilter("ignore")


class Template:
    constructed = False
    temp_output_dir = tempfile.TemporaryDirectory()

    def __init__(self, solvents, samples, diluent):
        self.wb = openpyxl.load_workbook(root_dir + "/template_file/HS_Quantification Template (HH v 2.0).xlsx")
        self.solvent_sheets = None
        self.solvents = solvents
        self.samples = samples
        self.diluent = diluent

        self.create_solvent_sheets()
        self.plot_chart()
        self.add_coa_data()
        self.add_area_height_data_a()
        self.add_area_height_data_b()
        self.add_sample_data()
        self.save_template()

    def create_solvent_sheets(self):
        """ Create a wb sheet for every solvent"""

        # Add a reference for each cell that is referenced in the original template:
        for j, solvent in enumerate(self.solvents):
            for cell in cells_with_reference[9:]:
                self.wb[f"solvent {j + 2}"][cell] = f"='{self.solvents[0].name}'!{cell}"
            self.wb[f"solvent {j + 1}"].title = solvent.name

        for cell in cells_with_reference[:9]:
            self.wb["Analytical Report"][cell] = f"='{self.solvents[0].name}'!{cell}"

        # Remove unused sheets (based on 12 + 1 solvent sheets):
        remaining_sheets = 13 - len(self.solvents)
        for sheet in range(remaining_sheets):
            self.wb.remove(self.wb[f"solvent {sheet + len(self.solvents) + 1}"])

        self.solvent_sheets = {solvent.name: self.wb[solvent.name] for solvent in self.solvents}

    def plot_chart(self):
        """ Generate a chart and specify the layout for all solvent sheets """

        for solvent in self.solvents:
            # Find x and y-values:
            xvalues = Reference(self.wb[solvent.name], min_col=5, min_row=62, max_row=69)
            yvalues = Reference(self.wb[solvent.name], min_col=7, min_row=62, max_row=69)

            # Instantiate chart object and set properties:
            chart = ScatterChart()
            chart.x_axis.title = "Concentration (µg/mL)"
            chart.x_axis.majorGridlines = None
            chart.x_axis.numFmt = '0'
            chart.y_axis.title = "Peak area (a.u.*s)"
            chart.y_axis.majorGridlines.spPr = GraphicalProperties(noFill='True')
            chart.y_axis.majorGridlines.spPr.ln = LineProperties(solidFill='C5E3F5')
            chart.legend = None

            # Instantiate series object and set properties:
            series = Series(yvalues, xvalues)
            series.marker = openpyxl.chart.marker.Marker('circle')
            series.graphicalProperties.line.noFill = True
            line_property = LineProperties(solidFill='A02D96')
            graphical_property = GraphicalProperties(ln=line_property)
            font_property = Font(typeface='Calibri Light')
            character_property = CharacterProperties(latin=font_property, sz=900, solidFill='FF0000')
            paragraph_property = ParagraphProperties(defRPr=character_property)
            richtext_property = RichText(p=[Paragraph(pPr=paragraph_property, endParaRPr=character_property)])
            number_format_property = NumFmt(formatCode='0.0000E+00')
            trendline_label = TrendlineLabel(txPr=richtext_property, numFmt=number_format_property)
            series.trendline = Trendline(dispRSqr=True, dispEq=True, spPr=graphical_property,
                                         trendlineLbl=trendline_label)
            chart.series.append(series)

            # Add chart to sheet.
            self.wb[solvent.name].add_chart(chart, "N61")

    def add_coa_data(self):

        # Add CoA data for 1. Diluent:
        try:
            self.solvent_sheets[self.solvents[0].name]["A22"] = self.diluent.name
            self.solvent_sheets[self.solvents[0].name]["B22"] = self.diluent.manufacturer
            self.solvent_sheets[self.solvents[0].name]["C22"] = self.diluent.catalog_number
            self.solvent_sheets[self.solvents[0].name]["D22"] = self.diluent.lot_number
        except Exception as e:
            print(e)

        # Add CoA for 2. Reference standards:
        for solvent in self.solvents:
            self.solvent_sheets[solvent.name]["A27"] = solvent.name
            self.solvent_sheets[solvent.name]["B27"] = solvent.manufacturer
            self.solvent_sheets[solvent.name]["C27"] = solvent.catalog_number
            self.solvent_sheets[solvent.name]["D27"] = solvent.lot_number
            self.solvent_sheets[solvent.name]["F27"] = solvent.expiration_date
            self.solvent_sheets[solvent.name]["E27"] = solvent.purity

    def add_area_height_data_a(self):

        # 6. Add peak area and peak height data for 6. Calibration curve:
        for solvent in self.solvents:

            for j in range(4):
                self.solvent_sheets[solvent.name][f"I{62 + j}"] = getattr(solvent, "a" + f"{8 - j}")[1]
                self.solvent_sheets[solvent.name][f"F{62 + j}"] = getattr(solvent, "a" + f"{8 - j}")[0]

            for j in range(4):
                self.solvent_sheets[solvent.name][f"F{66 + j}"] = getattr(solvent, "a" + f"{4 - j}")[0]

    def add_area_height_data_b(self):

        for solvent in self.solvents:

            # Add peak area and retention time data for 7. Repeatability and control:
            try:
                for j in range(3):
                    self.solvent_sheets[solvent.name][f"G{84 + j}"] = getattr(solvent, "b3_" + f"{j + 1}")[2]
                    self.solvent_sheets[solvent.name][f"H{84 + j}"] = getattr(solvent, "b3_" + f"{j + 1}")[0]
            except Exception as e:
                pass

            # # Add peak area for 8. Data for bracketing control:
            try:
                for j in range(5):
                    self.solvent_sheets[solvent.name][f"G{95 + j}"] = getattr(solvent, "b3_" + f"{j + 4}")[0]
            except Exception as e:
                print(e)

    def add_sample_data(self):

        for j in range(len(self.samples)):
            self.solvent_sheets[self.solvents[0].name][f"A{105 + (j * 6)}"] = self.samples[j].sample_code

        for solvent in self.solvents:
            for y, sample in enumerate(self.samples):
                for z in range(3):
                    self.solvent_sheets[solvent.name][f"I{105 + z + (y * 6)}"] = \
                        sample.__dict__[solvent.name][f"tag-{z + 1}"][0]
                    self.solvent_sheets[solvent.name][f"I{110 - z + (y * 6)}"] = \
                        sample.__dict__[solvent.name][f"tag-S-A{z + 4}"][0]

    def save_template(self):
        self.wb.save(Template.temp_output_dir.name + "/HS_Quantification Template (HH v 2.0) (processed).xlsx")
        Template.constructed = True